"""Paramètres du rapprochement (seuils) + Export Excel — tests backend.

Les valeurs de seuil utilisées ici (12.5 %, 20 L, 33 %) sont des valeurs
FIXTURE/TEST uniquement — JAMAIS des seuils métier de production.
Energy réel : NON TESTÉ (URL manquante) — mode par défaut non connecté.
"""
from __future__ import annotations

import io
import os

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL")
            or frontend_env.get("REACT_APP_BACKEND_URL")).rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@logitrak.ch", "password": "admin123"}
MANAGER = {"email": "manager@logitrak.ch", "password": "manager123"}
DRIVER = {"email": "chauffeur@logitrak.ch", "password": "chauffeur123"}
ADMIN_B = {"email": "admin-b@test.ch", "password": "testb123"}

SETTINGS_URL = f"{API}/livre/energy/reconciliation/settings"
EXPORT_URL = f"{API}/livre/energy/reconciliation/export.xlsx"
PARAMS_YEAR = {"date_from": "2026-01-01", "date_to": "2026-12-31"}


def _login(creds) -> dict:
    r = requests.post(f"{API}/auth/login", json=creds, timeout=20)
    assert r.status_code == 200, f"login {creds['email']}: {r.status_code}"
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def admin_h():
    return _login(ADMIN)


@pytest.fixture(scope="module")
def manager_h():
    return _login(MANAGER)


@pytest.fixture(scope="module")
def driver_h():
    return _login(DRIVER)


@pytest.fixture(scope="module")
def admin_b_h():
    return _login(ADMIN_B)


@pytest.fixture(scope="module", autouse=True)
def cleanup_thresholds(admin_h, admin_b_h):
    """Après les tests : remettre les seuils à null (aucune valeur par défaut métier)."""
    yield
    requests.put(SETTINGS_URL, json={"threshold_percent": None, "threshold_liters": None},
                 headers=admin_h, timeout=20)
    requests.put(SETTINGS_URL, json={"threshold_percent": None, "threshold_liters": None},
                 headers=admin_b_h, timeout=20)


# ---------------------------------------------------------------------------
# Paramètres — configuration, validation, RBAC, audit, isolation tenant
# ---------------------------------------------------------------------------
class TestReconciliationSettings:
    def test_default_no_threshold(self, admin_h):
        r = requests.get(SETTINGS_URL, headers=admin_h, timeout=20)
        assert r.status_code == 200
        b = r.json()
        assert b["alerting"] == "disabled"

    def test_save_and_read_threshold(self, admin_h):
        r = requests.put(SETTINGS_URL, json={"threshold_percent": 12.5,
                                             "threshold_liters": None},
                         headers=admin_h, timeout=20)  # 12.5 = valeur TEST/FIXTURE
        assert r.status_code == 200
        r = requests.get(SETTINGS_URL, headers=admin_h, timeout=20)
        assert r.json()["threshold_percent"] == 12.5
        assert r.json()["threshold_liters"] is None
        # préversion visible dans le preview
        r = requests.get(f"{API}/livre/energy/reconciliation/preview",
                         params=PARAMS_YEAR, headers=admin_h, timeout=90)
        th = r.json()["thresholds"]
        assert th["configured"] is True and th["percent"] == 12.5

    def test_save_liters_threshold(self, admin_h):
        r = requests.put(SETTINGS_URL, json={"threshold_percent": 12.5,
                                             "threshold_liters": 20},
                         headers=admin_h, timeout=20)  # 20 L = valeur TEST/FIXTURE
        assert r.status_code == 200
        assert requests.get(SETTINGS_URL, headers=admin_h, timeout=20).json()["threshold_liters"] == 20

    def test_invalid_values_rejected(self, admin_h):
        for payload in ({"threshold_percent": -5}, {"threshold_percent": 150},
                        {"threshold_percent": 0}, {"threshold_liters": -1},
                        {"threshold_liters": 0}):
            r = requests.put(SETTINGS_URL, json=payload, headers=admin_h, timeout=20)
            assert r.status_code == 400, f"{payload} aurait dû être rejeté"

    def test_rbac(self, manager_h, driver_h):
        # Manager : autorisé (convention PUT /livre/settings existante)
        r = requests.put(SETTINGS_URL, json={"threshold_percent": 12.5,
                                             "threshold_liters": 20},
                         headers=manager_h, timeout=20)
        assert r.status_code == 200
        # Chauffeur : jamais de modification du seuil flotte
        r = requests.put(SETTINGS_URL, json={"threshold_percent": 50},
                         headers=driver_h, timeout=20)
        assert r.status_code == 403
        r = requests.get(SETTINGS_URL, headers=driver_h, timeout=20)
        assert r.status_code == 403

    def test_audit_trail(self, admin_h):
        r = requests.get(f"{API}/livre/audit-log", params={"limit": 50},
                         headers=admin_h, timeout=20)
        assert r.status_code == 200
        entries = r.json()
        items = entries if isinstance(entries, list) else entries.get("items", [])
        hits = [e for e in items if e.get("action") == "energy.reconciliation_settings.update"]
        assert hits, "la modification du seuil doit être tracée dans l'audit"
        d = hits[0]["details"]
        assert "old" in d and "new" in d and d.get("tenant_id")

    def test_tenant_isolation(self, admin_h, admin_b_h):
        # Tenant B : aucun seuil hérité du tenant default
        r = requests.get(SETTINGS_URL, headers=admin_b_h, timeout=20)
        assert r.status_code == 200
        assert r.json()["threshold_percent"] is None
        # Tenant B configure son propre seuil (33 = TEST) sans affecter default
        requests.put(SETTINGS_URL, json={"threshold_percent": 33, "threshold_liters": None},
                     headers=admin_b_h, timeout=20)
        assert requests.get(SETTINGS_URL, headers=admin_h, timeout=20).json()["threshold_percent"] == 12.5
        assert requests.get(SETTINGS_URL, headers=admin_b_h, timeout=20).json()["threshold_percent"] == 33


# ---------------------------------------------------------------------------
# Statuts avec seuils (unitaires) — seuls les MEASURED exploitables → À contrôler
# ---------------------------------------------------------------------------
class TestThresholdStatusLogic:
    def _m(self, mt="MEASURED", value=100.0, availability="AVAILABLE"):
        return {"value": value, "unit": "L", "availability": availability,
                "measurement_type": mt, "source": "OBD", "timestamp": None}

    def _status(self, *a, **k):
        from app.routes.energy import _reconciliation_status
        return _reconciliation_status(*a, **k)

    def test_measured_beyond_liters_threshold(self):
        s, reason = self._status(True, 5, self._m(), 30.0, 3.0, None, 20)  # 20 L = TEST
        assert s == "A_CONTROLER" and "seuil 20 L" in reason

    def test_measured_within_liters_threshold(self):
        s, _ = self._status(True, 5, self._m(), 10.0, 3.0, None, 20)
        assert s == "OK"

    def test_negative_gap_absolute_value(self):
        s, _ = self._status(True, 5, self._m(), -30.0, -3.0, None, 20)
        assert s == "A_CONTROLER", "l'écart négatif doit être comparé en valeur absolue"

    def test_both_thresholds_any_exceeded(self):
        s, _ = self._status(True, 5, self._m(), 5.0, 25.0, 10, 50)
        assert s == "A_CONTROLER", "dépassement % suffit même si litres sous seuil"

    def test_estimated_never_a_controler(self):
        s, _ = self._status(True, 5, self._m(mt="ESTIMATED"), 500.0, 500.0, 1, 1)
        assert s == "INDICATIF"

    def test_reference_never_a_controler(self):
        s, _ = self._status(True, 5, self._m(mt="REFERENCE"), 500.0, 500.0, 1, 1)
        assert s == "INDICATIF"

    def test_stale_never_a_controler(self):
        s, _ = self._status(True, 5, self._m(availability="STALE"), 500.0, 500.0, 1, 1)
        assert s == "INDICATIF"

    def test_unmapped_never_a_controler(self):
        s, _ = self._status(False, 5, self._m(), 500.0, 500.0, 1, 1)
        assert s == "IMPOSSIBLE"

    def test_none_never_a_controler(self):
        s, _ = self._status(True, 5, None, None, None, 1, 1)
        assert s == "IMPOSSIBLE"


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------
def _export(headers, extra=None):
    params = {**PARAMS_YEAR, **(extra or {})}
    return requests.get(EXPORT_URL, params=params, headers=headers, timeout=120)


def _load_ws(resp):
    wb = load_workbook(io.BytesIO(resp.content))  # prouve que le fichier est ouvrable
    return wb.active


def _data_rows(ws):
    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "Véhicule":
            header_row = i
            headers = list(row)
            break
    assert header_row, "ligne d'entête introuvable"
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if any(v not in (None, "") for v in row):
            rows.append(dict(zip(headers, row)))
    return rows


class TestExport:
    def test_export_openable_and_complete(self, admin_h):
        r = _export(admin_h)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        assert "rapprochement_carburant_2026-01-01_2026-12-31.xlsx" in r.headers["content-disposition"]
        ws = _load_ws(r)
        rows = _data_rows(ws)
        assert len(rows) == 18, "les 18 véhicules (y compris non mappés) doivent être exportés"

    def test_metadata_header(self, admin_h):
        ws = _load_ws(_export(admin_h))
        head_text = " ".join(str(c[0]) for c in ws.iter_rows(max_row=8, max_col=1, values_only=True) if c[0])
        assert "Période : 2026-01-01 → 2026-12-31" in head_text
        assert "Alertes automatiques : désactivées" in head_text
        assert "Généré le" in head_text

    def test_null_never_zero(self, admin_h):
        rows = _data_rows(_load_ws(_export(admin_h)))
        for row in rows:
            if row.get("Type de mesure") == "Aucun":
                assert row.get("Consommation (L)") in (None, ""), \
                    "consommation absente → cellule vide, JAMAIS 0"
                assert row.get("Écart (L)") in (None, "")
                assert row.get("Écart (%)") in (None, "")

    def test_unmapped_vehicles_exported_with_reason(self, admin_h):
        rows = _data_rows(_load_ws(_export(admin_h)))
        unmapped = [r for r in rows if r.get("Tracker ID") in (None, "")]
        assert len(unmapped) == 6
        for r in unmapped:
            assert r.get("Statut") == "Impossible"
            assert "tracker Navixy" in (r.get("Raison du statut") or "")

    def test_french_labels(self, admin_h):
        rows = _data_rows(_load_ws(_export(admin_h)))
        for r in rows:
            assert r.get("Type de mesure") in ("Mesuré", "Estimé", "Référence", "Aucun")
            assert r.get("Statut") in ("OK", "À contrôler", "Indicatif", "Impossible")
            assert r.get("Fiabilité") in ("Exploitable", "Indicatif", "Impossible")

    def test_filter_vehicle(self, admin_h):
        preview = requests.get(f"{API}/livre/energy/reconciliation/preview",
                               params=PARAMS_YEAR, headers=admin_h, timeout=90).json()
        vid = preview["rows"][0]["vehicle_id"]
        rows = _data_rows(_load_ws(_export(admin_h, {"vehicle_id": vid})))
        assert len(rows) == 1

    def test_filter_status(self, admin_h):
        r = _export(admin_h, {"status": "IMPOSSIBLE"})
        rows = _data_rows(_load_ws(r))
        assert rows and all(row["Statut"] == "Impossible" for row in rows)

    def test_filter_measurement_none_disconnected(self, admin_h):
        r = _export(admin_h, {"measurement": "MEASURED"})
        assert r.status_code == 200
        rows = _data_rows(_load_ws(r))
        if requests.get(f"{API}/livre/energy/status", headers=admin_h,
                        timeout=20).json()["mode"] == "not_connected":
            assert rows == [], "non connecté → aucune consommation MEASURED"

    def test_rbac_driver_forbidden(self, driver_h):
        r = _export(driver_h)
        assert r.status_code == 403

    def test_tenant_isolation(self, admin_b_h, admin_h):
        rows_b = _data_rows(_load_ws(_export(admin_b_h))) if _export(admin_b_h).status_code == 200 else []
        plates_default = {row["Plaque"] for row in _data_rows(_load_ws(_export(admin_h)))}
        plates_b = {row.get("Plaque") for row in rows_b}
        assert not (plates_b & plates_default), "aucune plaque du tenant default dans l'export tenant B"

    def test_idor_vehicle_param(self, admin_b_h, admin_h):
        preview = requests.get(f"{API}/livre/energy/reconciliation/preview",
                               params=PARAMS_YEAR, headers=admin_h, timeout=90).json()
        vid = preview["rows"][0]["vehicle_id"]
        r = _export(admin_b_h, {"vehicle_id": vid})
        assert r.status_code == 200
        assert _data_rows(_load_ws(r)) == [], "IDOR : le véhicule du tenant default ne doit pas fuiter"

    def test_export_audited(self, admin_h):
        r = requests.get(f"{API}/livre/audit-log", params={"limit": 50},
                         headers=admin_h, timeout=20)
        entries = r.json()
        items = entries if isinstance(entries, list) else entries.get("items", [])
        assert any(e.get("action") == "energy.reconciliation.export" for e in items)
