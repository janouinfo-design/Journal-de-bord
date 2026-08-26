"""GARDE-FOU BEV — legacy 0,085 L/km (T1–T15, T11 UI couvert par testing agent).

Règle : le legacy fuel_l = distance_km × 0.085 n'est écrit que pour ICE
(ice/diesel/essence/petrol) et UNKNOWN (dette résiduelle documentée —
comportement historique conservé, jamais converti en BEV/ICE).
BEV / HEV / PHEV → fuel_l ABSENT (jamais 0 — zéro signifierait une mesure).
AUCUNE migration Mongo : l'historique reste intact.
"""
import asyncio
import io
import os
import uuid

import pymongo
import pytest
import requests

from app import navixy_sync as ns
from app import reports as rp

API = os.environ["REACT_APP_BACKEND_URL"].rstrip("/") + "/api"


def _login(email, password):
    r = requests.post(f"{API}/auth/login",
                      json={"email": email, "password": password}, timeout=15)
    assert r.status_code == 200, f"login {email}: {r.status_code}"
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def admin_h():
    return _login("admin@logitrak.ch", "admin123")


@pytest.fixture(scope="module")
def mongo():
    mc = pymongo.MongoClient(os.environ["MONGO_URL"], serverSelectionTimeoutMS=5000)
    yield mc[os.environ["DB_NAME"]]
    mc.close()


# ---------------------------------------------------------------------------
# T1–T8 (unitaires) — helper centralisé legacy_fuel_estimation_allowed
# ---------------------------------------------------------------------------
class TestHelper:
    def test_t1_ice(self):
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "ice"}) is True

    def test_t2_diesel(self):
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "diesel"}) is True
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "Diesel"}) is True

    def test_t3_essence(self):
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "essence"}) is True
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "petrol"}) is True

    def test_t4_bev_blocked(self):
        for ft in ("electric", "bev", "ev", "ELECTRIC"):
            assert ns.legacy_fuel_estimation_allowed({"fuel_type": ft}) is False

    def test_t6_unknown_stays_unknown(self):
        # UNKNOWN n'est converti ni en BEV ni en ICE : le mapping reste UNKNOWN,
        # et le legacy historique reste actif (dette résiduelle documentée).
        for ft in (None, "", "hydrogen", "gaz"):
            assert ns.powertrain_from_fuel_type(ft) == "UNKNOWN"
            assert ns.legacy_fuel_estimation_allowed({"fuel_type": ft}) is True
        assert ns.legacy_fuel_estimation_allowed({}) is True
        assert ns.legacy_fuel_estimation_allowed(None) is True

    def test_t7_hev_ambiguous_no_legacy(self):
        # Aucune convention 8,5 L/100 validée pour hybrides → pas de calcul.
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "hybrid"}) is False
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "hev"}) is False

    def test_t8_phev_ambiguous_no_legacy(self):
        assert ns.legacy_fuel_estimation_allowed({"fuel_type": "phev"}) is False

    def test_mapping_canonical_not_duplicated(self):
        # routes/energy réutilise le mapping canonique (pas de duplication).
        from app.routes.energy import _powertrain
        assert _powertrain("electric") == "BEV"
        assert _powertrain("diesel") == "ICE"
        assert _powertrain(None) == "UNKNOWN"


# ---------------------------------------------------------------------------
# T4/T5/T6 (niveau document) — _build_trip_doc n'écrit fuel_l que si autorisé
# ---------------------------------------------------------------------------
def _build_doc(vehicle, monkeypatch, km=100.0):
    async def _no_driver(db, vid, start):
        return None
    monkeypatch.setattr(ns, "resolve_driver_for_trip", _no_driver)
    tr = {"id": f"gf-{uuid.uuid4()}", "start_date": "2026-08-26 08:00:00",
          "end_date": "2026-08-26 09:00:00", "length": km,
          "bounds": {}, "avg_speed": 40, "max_speed": 80}
    return asyncio.run(ns._build_trip_doc(None, vehicle, tr, []))


class TestBuildTripDoc:
    VEH = {"id": "veh-gf", "plate": "GF 000", "fuel_type": None}

    def test_t4_bev_no_fuel_l(self, monkeypatch):
        doc = _build_doc({**self.VEH, "fuel_type": "electric"}, monkeypatch)
        assert "fuel_l" not in doc, "BEV : fuel_l doit être ABSENT"

    def test_t5_bev_never_zero(self, monkeypatch):
        doc = _build_doc({**self.VEH, "fuel_type": "electric"}, monkeypatch)
        assert doc.get("fuel_l") != 0  # absent ≠ 0 mesuré

    def test_ice_legacy_computed(self, monkeypatch):
        doc = _build_doc({**self.VEH, "fuel_type": "diesel"}, monkeypatch, km=100.0)
        assert doc["fuel_l"] == round(100.0 * ns.FUEL_L_PER_KM, 2) == 8.5

    def test_t6_unknown_keeps_historical_behaviour(self, monkeypatch):
        doc = _build_doc(self.VEH, monkeypatch, km=100.0)
        assert doc["fuel_l"] == 8.5  # dette résiduelle documentée

    def test_t7_t8_hybrids_no_fuel_l(self, monkeypatch):
        for ft in ("hybrid", "phev"):
            doc = _build_doc({**self.VEH, "fuel_type": ft}, monkeypatch)
            assert "fuel_l" not in doc, f"{ft} : cas ambigu → pas de calcul legacy"


# ---------------------------------------------------------------------------
# T9 — l'historique Mongo n'est PAS modifié (aucune migration)
# ---------------------------------------------------------------------------
class TestHistoricalUntouched:
    def test_t9_pilot_trip_keeps_legacy_fuel(self, mongo):
        pilot = mongo.trips.find_one(
            {"id": "8be7b16f-ff74-4262-a1ff-1735f085c3d4"}, {"_id": 0, "fuel_l": 1})
        assert pilot is not None
        assert pilot.get("fuel_l") == 0.66  # valeur legacy historique intacte

    def test_t9_historical_volume_still_has_fuel(self, mongo):
        n = mongo.trips.count_documents({"fuel_l": {"$exists": True}})
        assert n >= 5000, "les fuel_l legacy historiques doivent rester persistés"


# ---------------------------------------------------------------------------
# Unitaires exports — un trajet sans fuel_l ne produit jamais 0
# ---------------------------------------------------------------------------
_TRIP_BEV = {
    "start_time": "2023-05-02T08:00:00+00:00", "end_time": "2023-05-02T09:00:00+00:00",
    "driver_name": "Test GF", "vehicle_plate": "GF BEV 001",
    "start_address": "A", "end_address": "B",
    "distance_km": 100.0, "duration_min": 60, "avg_speed": 40, "max_speed": 80,
}
_TRIP_ICE = {**_TRIP_BEV, "vehicle_plate": "GF ICE 001", "fuel_l": 8.5}


class TestUnitExports:
    def test_csv_empty_cell_not_zero(self):
        out = rp.trips_to_csv([_TRIP_BEV, _TRIP_ICE], "Professionnel").decode("utf-8-sig")
        bev_line = next(l for l in out.splitlines() if "GF BEV 001" in l)
        ice_line = next(l for l in out.splitlines() if "GF ICE 001" in l)
        assert bev_line.split(";")[8] == "", "CSV BEV : cellule carburant vide, jamais 0"
        assert ice_line.split(";")[8] == "8.5"

    def test_xlsx_empty_cell_not_zero(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(rp.trips_to_xlsx([_TRIP_BEV, _TRIP_ICE],
                                                       "Professionnel", "Test GF")))
        rows = list(wb.active.iter_rows(values_only=True))
        bev = next(r for r in rows if r and "GF BEV 001" in r)
        ice = next(r for r in rows if r and "GF ICE 001" in r)
        assert bev[8] is None or bev[8] == "", "XLSX BEV : cellule vide, jamais 0"
        assert ice[8] == 8.5

    def test_t13_pdf_dash_not_zero(self):
        import fitz
        pdf = rp.trips_to_pdf([_TRIP_BEV, _TRIP_ICE], "Professionnel", "Test GF")
        text = "".join(p.get_text() for p in fitz.open(stream=pdf, filetype="pdf"))
        assert "—" in text, "PDF : trajet sans fuel_l → tiret"
        assert "0.00" not in text, "PDF : jamais 0.00 L fabriqué pour un BEV"
        assert "8.50" in text  # témoin thermique + total honnête


# ---------------------------------------------------------------------------
# T10/T12/T13/T14/T15 — E2E : véhicule BEV + trajets test (année 2023 isolée)
# ---------------------------------------------------------------------------
@pytest.fixture(scope="class")
def seeded(mongo):
    """Véhicule BEV + véhicule diesel témoin + 1 trajet chacun (2023, isolé)."""
    tag = uuid.uuid4().hex[:8]
    veh_bev = {"id": f"gf-bev-{tag}", "tenant_id": "default",
               "plate": f"GF-BEV-{tag}", "label": "Test garde-fou BEV",
               "fuel_type": "electric"}
    veh_ice = {"id": f"gf-ice-{tag}", "tenant_id": "default",
               "plate": f"GF-ICE-{tag}", "label": "Test garde-fou ICE",
               "fuel_type": "diesel"}
    base = {"tenant_id": "default", "driver_id": None,
            "start_time": "2023-05-02T08:00:00+00:00",
            "end_time": "2023-05-02T09:00:00+00:00",
            "start_address": "A", "end_address": "B",
            "distance_km": 100.0, "duration_min": 60,
            "avg_speed": 40.0, "max_speed": 80.0,
            "classification": "professional", "auto_classified": True}
    trip_bev = {**base, "id": f"gf-trip-bev-{tag}", "vehicle_id": veh_bev["id"],
                "vehicle_plate": veh_bev["plate"], "driver_name": veh_bev["plate"],
                "navixy_track_id": f"gf-{tag}-bev"}  # PAS de fuel_l (nouveau garde-fou)
    trip_ice = {**base, "id": f"gf-trip-ice-{tag}", "vehicle_id": veh_ice["id"],
                "vehicle_plate": veh_ice["plate"], "driver_name": veh_ice["plate"],
                "navixy_track_id": f"gf-{tag}-ice", "fuel_l": 8.5}
    mongo.vehicles.insert_many([dict(veh_bev), dict(veh_ice)])
    mongo.trips.insert_many([dict(trip_bev), dict(trip_ice)])
    yield {"veh_bev": veh_bev, "veh_ice": veh_ice,
           "trip_bev": trip_bev, "trip_ice": trip_ice}
    mongo.trips.delete_many({"id": {"$in": [trip_bev["id"], trip_ice["id"]]}})
    mongo.vehicles.delete_many({"id": {"$in": [veh_bev["id"], veh_ice["id"]]}})


class TestApiAndExportsE2E:
    def test_t10_api_trips_bev_no_fuel_l(self, admin_h, seeded):
        r = requests.get(f"{API}/livre/trips",
                         params={"vehicle_id": seeded["veh_bev"]["id"],
                                 "start": "2023-01-01", "end": "2023-12-31"},
                         headers=admin_h, timeout=20)
        assert r.status_code == 200
        trips = r.json()["trips"]
        rows = [t for t in trips if t["id"] == seeded["trip_bev"]["id"]]
        assert rows, "trajet BEV test introuvable via l'API"
        t = rows[0]
        assert t.get("fuel_l") is None, "T15 : absence sérialisée null/absent, jamais 0"
        assert "fuel_l" not in t or t["fuel_l"] is None

    def test_t12_export_xlsx_bev_empty(self, admin_h, seeded):
        from openpyxl import load_workbook
        r = requests.get(f"{API}/livre/reports/export",
                         params={"classification": "professional", "fmt": "xlsx",
                                 "start": "2023-01-01", "end": "2023-12-31"},
                         headers=admin_h, timeout=60)
        assert r.status_code == 200
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        bev = next(r_ for r_ in rows if r_ and seeded["veh_bev"]["plate"] in r_)
        ice = next(r_ for r_ in rows if r_ and seeded["veh_ice"]["plate"] in r_)
        assert bev[8] in (None, ""), "export BEV : Carburant estimé (L) vide, jamais 0/8.5"
        assert ice[8] == 8.5

    def test_t13_pdf_export_and_tax(self, admin_h, seeded):
        import fitz
        r = requests.get(f"{API}/livre/reports/export",
                         params={"classification": "professional", "fmt": "pdf",
                                 "start": "2023-01-01", "end": "2023-12-31"},
                         headers=admin_h, timeout=60)
        assert r.status_code == 200
        text = "".join(p.get_text() for p in fitz.open(stream=r.content, filetype="pdf"))
        assert "—" in text and "0.00" not in text
        # PDF fiscal 2023 : seul le témoin diesel contribue des litres estimés.
        r2 = requests.get(f"{API}/livre/reports/tax-swiss", params={"year": 2023},
                          headers=admin_h, timeout=60)
        assert r2.status_code == 200
        tax = "".join(p.get_text() for p in fitz.open(stream=r2.content, filetype="pdf"))
        assert "Estimé*" in tax
        assert "8.50" in tax, "total fiscal = 8.50 L (diesel) — le BEV n'ajoute AUCUN litre"

    def test_t14_trip_energy_block_independent(self, admin_h, seeded):
        """TripEnergyBlock (batch Energy) reste indépendant du legacy fuel_l."""
        r = requests.post(f"{API}/livre/energy/trips",
                          json={"trip_ids": [seeded["trip_bev"]["id"]]},
                          headers=admin_h, timeout=60)
        assert r.status_code == 200
        env = r.json()["results"][0]
        assert env["trip_id"] == seeded["trip_bev"]["id"]
        # véhicule test sans tracker → indisponibilité honnête, jamais le legacy en fallback
        assert env["availability"] == "UNAVAILABLE"
        fuel = env.get("fuel")
        if fuel:
            for m in fuel.values():
                if isinstance(m, dict):
                    assert m.get("value") is None, "jamais le legacy 8.5 injecté dans Energy"
