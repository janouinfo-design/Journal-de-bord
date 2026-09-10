"""Rapport mensuel fiscal chauffeur — tests backend.

SMTP MOCKÉ UNIQUEMENT : aucun email réel ne doit être envoyé. Les tests du job
patchent `app.emailer.is_smtp_configured` et `app.emailer.send_email_with_attachments`.
Aucun appel Navixy / device-write n'est impliqué par ce module.
Données seedées : trips `TEST-MKR-*` sur mars 2025 (+ février pour le cumul annuel),
nettoyées en fin de module. Config settings sauvegardée/restaurée.
"""
from __future__ import annotations

import asyncio
import io
import os
from datetime import datetime

import pytest
import pymongo
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL")
            or frontend_env.get("REACT_APP_BACKEND_URL")).rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@logitrak.ch", "password": "admin123"}
MANAGER = {"email": "manager@logitrak.ch", "password": "manager123"}
DRIVER = {"email": "chauffeur@logitrak.ch", "password": "chauffeur123"}
ADMIN_B = {"email": "admin-b@test.ch", "password": "testb123"}

CONFIG_URL = f"{API}/livre/reports/monthly-driver/config"
DOWNLOAD_URL = f"{API}/livre/reports/monthly-driver/download"

# Mois de test : mars 2025 — CET au début (UTC+1), CEST à la fin (UTC+2, DST le 30/03).
# Bornes UTC attendues : [2025-02-28T23:00, 2025-03-31T22:00).
TEST_MONTH = "2025-03"
MARK = "TEST-MKR"

SEED_TRIPS = [
    # Alice mars : 100.5 pro + 50.25 privé + 10.0 non classifié = 160.75 ; % privé = 33.3
    {"id": f"{MARK}-T1", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-03T08:00:00+00:00", "distance_km": 100.5, "classification": "professional"},
    {"id": f"{MARK}-T2", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-15T18:00:00+00:00", "distance_km": 50.25, "classification": "personal"},
    {"id": f"{MARK}-T3", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-03-20T10:00:00+00:00", "distance_km": 10.0, "classification": None},
    # Alice février (cumul annuel uniquement) : +40 pro
    {"id": f"{MARK}-T6", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-02-10T08:00:00+00:00", "distance_km": 40.0, "classification": "professional"},
    # Alice HORS mois : 1er avril 00:30 CEST (2025-03-31T22:30 UTC) — ni mois ni cumul
    {"id": f"{MARK}-T5", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-31T22:30:00+00:00", "distance_km": 999.0, "classification": "professional"},
    # Bob : 31 mars 23:30 CEST (21:30 UTC) — DANS le mois
    {"id": f"{MARK}-T4", "driver_id": f"{MARK}-D2", "driver_name": f"{MARK} Bob",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-03-31T21:30:00+00:00", "distance_km": 200.0, "classification": "professional"},
    # Bob : 1er mars 00:30 CET (2025-02-28T23:30 UTC) — DANS le mois
    {"id": f"{MARK}-T8", "driver_id": f"{MARK}-D2", "driver_name": f"{MARK} Bob",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-02-28T23:30:00+00:00", "distance_km": 30.0, "classification": "professional"},
    # Bob : 28 février 23:30 CET (22:30 UTC) — HORS mois, DANS cumul annuel
    {"id": f"{MARK}-T9", "driver_id": f"{MARK}-D2", "driver_name": f"{MARK} Bob",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-02-28T22:30:00+00:00", "distance_km": 7.0, "classification": "professional"},
    # Carol : uniquement non classifié → % privé non calculable (null, jamais 0)
    {"id": f"{MARK}-T7", "driver_id": f"{MARK}-D3", "driver_name": f"{MARK} Carol",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-10T10:00:00+00:00", "distance_km": 5.0, "classification": None},
]


def _login(creds) -> dict:
    r = requests.post(f"{API}/auth/login", json=creds, timeout=20)
    assert r.status_code == 200, f"login {creds['email']}: {r.status_code}"
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def admin_h():
    return _login(ADMIN)


@pytest.fixture(scope="module")
def manager_h():
    return _login(MANAGER)


@pytest.fixture(scope="module")
def driver_h():
    return _login(DRIVER)


@pytest.fixture(scope="module")
def admin_b_h():
    return _login(ADMIN_B)


@pytest.fixture(scope="module")
def mongo():
    mc = pymongo.MongoClient(os.environ["MONGO_URL"], serverSelectionTimeoutMS=5000)
    yield mc[os.environ["DB_NAME"]]
    mc.close()


@pytest.fixture(scope="module", autouse=True)
def seed_and_restore(mongo):
    settings_q = {"tenant_id": "default", "id": "default"}
    original = mongo.settings.find_one(settings_q, {"_id": 0,
                                                    "monthly_report_enabled": 1,
                                                    "monthly_report_recipients": 1}) or {}
    mongo.trips.delete_many({"id": {"$regex": f"^{MARK}-"}})
    now = datetime.utcnow().isoformat() + "+00:00"
    for t in SEED_TRIPS:
        mongo.trips.insert_one({**t, "tenant_id": "default", "created_at": now,
                                "end_time": t["start_time"], "duration_min": 30,
                                "auto_classified": False})
    yield
    mongo.trips.delete_many({"id": {"$regex": f"^{MARK}-"}})
    mongo.monthly_report_log.delete_many({"month": TEST_MONTH})
    restore = {"monthly_report_enabled": original.get("monthly_report_enabled", False),
               "monthly_report_recipients": original.get("monthly_report_recipients", [])}
    mongo.settings.update_one(settings_q, {"$set": restore})


def _xlsx_rows(content: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _find_row(rows, name):
    for r in rows:
        if r and r[0] == name:
            return r
    return None


# ---------------------------------------------------------------------------
# 1. Fonctions pures de période
# ---------------------------------------------------------------------------
def test_period_functions():
    from zoneinfo import ZoneInfo
    from app.monthly_driver_report import (month_bounds, month_key, month_label,
                                           parse_month, previous_month, year_start_iso)
    assert parse_month("2025-03") == (2025, 3)
    for bad in ("2025-13", "2025-00", "pas-un-mois", "2025-3", ""):
        with pytest.raises(ValueError):
            parse_month(bad)
    assert month_key(2025, 3) == "2025-03"
    assert month_label(2025, 3) == "mars 2025"
    # Le 1er avril 06:00 Zurich → mois précédent = mars ; le 1er janvier → décembre N-1
    assert previous_month(datetime(2025, 4, 1, 6, 0, tzinfo=ZoneInfo("Europe/Zurich"))) == (2025, 3)
    assert previous_month(datetime(2025, 1, 1, 6, 0, tzinfo=ZoneInfo("Europe/Zurich"))) == (2024, 12)
    # Mars 2025 : CET au début, CEST à la fin (DST 30/03)
    start, end = month_bounds(2025, 3)
    assert start == "2025-02-28T23:00:00+00:00"
    assert end == "2025-03-31T22:00:00+00:00"
    # Décembre → 1er janvier de l'année suivante
    assert month_bounds(2025, 12)[1] == "2025-12-31T23:00:00+00:00"
    assert year_start_iso(2025) == "2024-12-31T23:00:00+00:00"


# ---------------------------------------------------------------------------
# 2. Configuration — GET/PUT + validation + RBAC
# ---------------------------------------------------------------------------
def test_config_get_defaults(admin_h):
    r = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r.status_code == 200, r.text
    d = r.json()
    assert isinstance(d["enabled"], bool)
    assert isinstance(d["recipients"], list)
    assert d["formats"] == ["pdf", "xlsx"]
    assert "smtp_configured" in d
    assert "Mois civil précédent" in d["period_rule"]
    assert d["send_day"] == "1er du mois 06:00 (Europe/Zurich)"
    assert len(d["default_month"]) == 7  # YYYY-MM


def test_config_put_and_normalisation(admin_h):
    r = requests.put(CONFIG_URL, headers=admin_h, timeout=20, json={
        "enabled": True,
        "recipients": ["  Fiduciaire@Logitrak.CH ", "fiduciaire@logitrak.ch", "rh@logitrak.ch"]})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["enabled"] is True
    assert d["recipients"] == ["fiduciaire@logitrak.ch", "rh@logitrak.ch"]
    r2 = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r2.json()["recipients"] == ["fiduciaire@logitrak.ch", "rh@logitrak.ch"]


def test_config_independent_from_weekly(admin_h, mongo):
    """La config mensuelle ne touche pas aux clés hebdomadaires."""
    doc = mongo.settings.find_one({"tenant_id": "default", "id": "default"}, {"_id": 0})
    assert doc.get("monthly_report_recipients") == ["fiduciaire@logitrak.ch", "rh@logitrak.ch"]
    assert doc.get("weekly_report_recipients", []) != ["fiduciaire@logitrak.ch", "rh@logitrak.ch"]


def test_config_validation(admin_h):
    r = requests.put(CONFIG_URL, headers=admin_h, timeout=20,
                     json={"recipients": ["pas-un-email"]})
    assert r.status_code == 400
    r2 = requests.put(CONFIG_URL, headers=admin_h, timeout=20,
                      json={"recipients": [f"u{i}@x.ch" for i in range(21)]})
    assert r2.status_code == 400


def test_config_rbac(manager_h, driver_h):
    assert requests.get(CONFIG_URL, headers=manager_h, timeout=20).status_code == 200
    assert requests.put(CONFIG_URL, headers=manager_h, timeout=20,
                        json={"enabled": True}).status_code == 200
    assert requests.get(CONFIG_URL, headers=driver_h, timeout=20).status_code == 403
    assert requests.put(CONFIG_URL, headers=driver_h, timeout=20,
                        json={"enabled": True}).status_code == 403
    assert requests.get(DOWNLOAD_URL, headers=driver_h, timeout=20,
                        params={"month": TEST_MONTH, "format": "pdf"}).status_code == 403


# ---------------------------------------------------------------------------
# 3. Téléchargement manuel — PDF / XLSX / mois vide / isolation
# ---------------------------------------------------------------------------
def test_download_pdf(admin_h):
    r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                     params={"month": TEST_MONTH, "format": "pdf"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert r.content[:4] == b"%PDF"
    assert f"rapport_mensuel_fiscal_chauffeurs_{TEST_MONTH}.pdf" in r.headers.get("content-disposition", "")


def test_download_xlsx_values(admin_h):
    r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                     params={"month": TEST_MONTH, "format": "xlsx"})
    assert r.status_code == 200, r.text
    rows = _xlsx_rows(r.content)
    # Colonnes : Chauffeur, Km PRO, Km PRIVÉ, Km non classifiés, Km total, % privé,
    #   Trajets, Trajets non classifiés, Véhicules, Cumul PRO, Cumul PRIVÉ, Statut fiscal
    alice = _find_row(rows, f"{MARK} Alice")
    assert alice, "Ligne Alice absente du XLSX"
    assert alice[1] == 100.5          # T5 (999 km, 1er avril 00:30 CEST) exclu
    assert alice[2] == 50.25
    assert alice[3] == 10.0
    assert alice[4] == 160.75         # total = TOUS les km
    assert alice[5] == 33.3           # 50.25 / 150.75
    assert alice[6] == 3 and alice[7] == 1
    assert f"{MARK} 111" in alice[8] and f"{MARK} 222" in alice[8]
    assert alice[9] == 140.5          # cumul annuel PRO = 100.5 + 40 (février)
    assert alice[10] == 50.25
    assert alice[11] == "À compléter (1)"

    bob = _find_row(rows, f"{MARK} Bob")
    assert bob, "Ligne Bob absente"
    assert bob[1] == 230.0            # 200 (31/03 23:30 CEST) + 30 (01/03 00:30 CET)
    assert bob[5] == 0.0              # 0 km privé sur des km classifiés = vrai zéro
    assert bob[6] == 2
    assert bob[9] == 237.0            # + 7 km du 28 février (cumul)
    assert bob[11] == "Complet"

    carol = _find_row(rows, f"{MARK} Carol")
    assert carol, "Ligne Carol absente"
    assert carol[5] is None           # aucun km classifié → % privé null, JAMAIS 0
    assert carol[11] == "À compléter (1)"

    total = _find_row(rows, "TOTAL")
    assert total and total[11] == "2 chauffeur(s) à compléter"


def test_download_empty_month_is_valid_report(admin_h):
    for fmt in ("pdf", "xlsx"):
        r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                         params={"month": "2020-01", "format": fmt})
        assert r.status_code == 200, f"{fmt}: {r.status_code}"
    rows = _xlsx_rows(r.content)
    assert any(row and row[0] == "Aucun trajet sur la période" for row in rows)


def test_download_bad_params(admin_h):
    for bad in ("pas-un-mois", "2025-13", "2025-03-01"):
        assert requests.get(DOWNLOAD_URL, headers=admin_h, timeout=20,
                            params={"month": bad, "format": "pdf"}).status_code == 400, bad
    assert requests.get(DOWNLOAD_URL, headers=admin_h, timeout=20,
                        params={"month": TEST_MONTH, "format": "docx"}).status_code == 422


def test_tenant_isolation(admin_b_h):
    r = requests.get(CONFIG_URL, headers=admin_b_h, timeout=20)
    assert r.status_code == 200
    assert r.json()["recipients"] != ["fiduciaire@logitrak.ch", "rh@logitrak.ch"] or \
        r.json()["recipients"] == []
    rx = requests.get(DOWNLOAD_URL, headers=admin_b_h, timeout=30,
                      params={"month": TEST_MONTH, "format": "xlsx"})
    assert rx.status_code == 200
    rows = _xlsx_rows(rx.content)
    assert _find_row(rows, f"{MARK} Alice") is None
    assert _find_row(rows, f"{MARK} Bob") is None


# ---------------------------------------------------------------------------
# 4. Job scheduler — SMTP mocké, idempotence, statuts
# ---------------------------------------------------------------------------
def _run_job(monkeypatch, mongo, *, smtp_configured, send_ok=True,
             enabled=True, recipients=None):
    import app.emailer as emailer_mod

    if recipients is None:
        recipients = ["fiduciaire@logitrak.ch"]
    mongo.settings.update_one(
        {"tenant_id": "default", "id": "default"},
        {"$set": {"monthly_report_enabled": enabled,
                  "monthly_report_recipients": recipients}})

    sent = []

    async def fake_send(to_list, subject, text, html, attachments):
        sent.append({"to": list(to_list), "subject": subject, "text": text,
                     "attachments": [(a[0], len(a[1])) for a in attachments]})
        return send_ok

    monkeypatch.setattr(emailer_mod, "is_smtp_configured", lambda: smtp_configured)
    monkeypatch.setattr(emailer_mod, "send_email_with_attachments", fake_send)

    async def runner():
        from app.db import close_db, init_db
        from app.monthly_driver_report import run_monthly_report_job
        init_db()
        try:
            return await run_monthly_report_job(year_month=TEST_MONTH)
        finally:
            close_db()

    results = asyncio.run(runner())
    return results, sent


@pytest.fixture()
def clean_log(mongo):
    mongo.monthly_report_log.delete_many({"month": TEST_MONTH})
    yield
    mongo.monthly_report_log.delete_many({"month": TEST_MONTH})


def test_job_sends_one_email_with_two_attachments(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True)
    assert results.get("default") == "sent"
    assert len(sent) == 1
    mail = sent[0]
    assert mail["to"] == ["fiduciaire@logitrak.ch"]
    assert "mars 2025" in mail["subject"]
    assert "Chauffeurs à compléter avant déclaration : 2" in mail["text"]
    names = [a[0] for a in mail["attachments"]]
    assert names == [f"rapport_mensuel_fiscal_chauffeurs_{TEST_MONTH}.pdf",
                     f"rapport_mensuel_fiscal_chauffeurs_{TEST_MONTH}.xlsx"]
    assert all(a[1] > 1000 for a in mail["attachments"])
    log = mongo.monthly_report_log.find_one({"tenant_id": "default", "month": TEST_MONTH})
    assert log and log["status"] == "sent" and log["error"] is None

    # Idempotence : relance même mois -> aucun nouvel envoi
    results2, sent2 = _run_job(monkeypatch, mongo, smtp_configured=True)
    assert results2.get("default") == "already_sent"
    assert len(sent2) == 0
    assert mongo.monthly_report_log.count_documents(
        {"tenant_id": "default", "month": TEST_MONTH, "status": "sent"}) == 1
    # Le journal hebdomadaire n'est pas touché
    assert mongo.weekly_report_log.count_documents({"week_monday": TEST_MONTH}) == 0


def test_job_smtp_not_configured(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=False)
    assert results.get("default") == "not_sent_smtp_not_configured"
    assert sent == []
    log = mongo.monthly_report_log.find_one({"tenant_id": "default", "month": TEST_MONTH})
    assert log["status"] == "not_sent_smtp_not_configured"
    results2, _ = _run_job(monkeypatch, mongo, smtp_configured=False)
    assert results2.get("default") == "not_sent_smtp_not_configured"


def test_job_no_recipients(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, recipients=[])
    assert results.get("default") == "not_sent_no_recipients"
    assert sent == []


def test_job_disabled(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, enabled=False)
    assert "default" not in results
    assert sent == []
    assert mongo.monthly_report_log.count_documents(
        {"tenant_id": "default", "month": TEST_MONTH}) == 0


def test_job_smtp_error(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, send_ok=False)
    assert results.get("default") == "error"
    assert len(sent) == 1
    log = mongo.monthly_report_log.find_one({"tenant_id": "default", "month": TEST_MONTH})
    assert log["status"] == "error" and log["error"]


def test_config_reports_last_send(admin_h, monkeypatch, mongo, clean_log):
    _run_job(monkeypatch, mongo, smtp_configured=True)
    r = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r.status_code == 200
    last = r.json()["last_send"]
    assert last and last["month"] == TEST_MONTH and last["status"] == "sent"
