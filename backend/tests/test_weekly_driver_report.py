"""Rapport hebdomadaire chauffeur — tests backend.

SMTP MOCKÉ UNIQUEMENT : aucun email réel ne doit être envoyé. Les tests du job
patchent `app.emailer.is_smtp_configured` et `app.emailer.send_email_with_attachments`.
Aucun appel Navixy / device-write n'est impliqué par ce module.
Données seedées : trips `TEST-WKR-*` sur la semaine du 2025-03-03 (hors données réelles),
nettoyées en fin de module. Config settings sauvegardée/restaurée.
"""
from __future__ import annotations

import asyncio
import io
import os
from datetime import date, datetime

import pytest
import pymongo
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL")
            or frontend_env.get("REACT_APP_BACKEND_URL")).rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "admin@logitrak.ch", "password": "admin123"}
MANAGER = {"email": "manager@logitrak.ch", "password": "manager123"}
DRIVER = {"email": "chauffeur@logitrak.ch", "password": "chauffeur123"}
ADMIN_B = {"email": "admin-b@test.ch", "password": "testb123"}

CONFIG_URL = f"{API}/livre/reports/weekly-driver/config"
DOWNLOAD_URL = f"{API}/livre/reports/weekly-driver/download"

# Semaine de test : lundi 2025-03-03 -> dimanche 2025-03-09 (CET = UTC+1).
# Bornes UTC attendues : [2025-03-02T23:00, 2025-03-09T23:00).
TEST_WEEK = "2025-03-03"
MARK = "TEST-WKR"

SEED_TRIPS = [
    # Alice : 100.5 pro + 50.25 privé + 10.0 non classifié = 160.75 total, 3 trajets
    {"id": f"{MARK}-T1", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-03T08:00:00+00:00", "end_time": "2025-03-03T09:00:00+00:00",
     "distance_km": 100.5, "classification": "professional"},
    {"id": f"{MARK}-T2", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-05T18:00:00+00:00", "end_time": "2025-03-05T19:00:00+00:00",
     "distance_km": 50.25, "classification": "personal"},
    {"id": f"{MARK}-T3", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-03-06T10:00:00+00:00", "end_time": "2025-03-06T10:30:00+00:00",
     "distance_km": 10.0, "classification": None},
    # Bob : borne dimanche 23:30 heure locale (22:30 UTC) — DANS la semaine
    {"id": f"{MARK}-T4", "driver_id": f"{MARK}-D2", "driver_name": f"{MARK} Bob",
     "vehicle_id": f"{MARK}-V2", "vehicle_plate": f"{MARK} 222",
     "start_time": "2025-03-09T22:30:00+00:00", "end_time": "2025-03-09T22:50:00+00:00",
     "distance_km": 200.0, "classification": "professional"},
    # HORS semaine : lundi suivant 00:30 heure locale (2025-03-09T23:30 UTC)
    {"id": f"{MARK}-T5", "driver_id": f"{MARK}-D1", "driver_name": f"{MARK} Alice",
     "vehicle_id": f"{MARK}-V1", "vehicle_plate": f"{MARK} 111",
     "start_time": "2025-03-09T23:30:00+00:00", "end_time": "2025-03-10T00:10:00+00:00",
     "distance_km": 999.0, "classification": "professional"},
]


def _login(creds) -> dict:
    r = requests.post(f"{API}/auth/login", json=creds, timeout=20)
    assert r.status_code == 200, f"login {creds['email']}: {r.status_code}"
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def admin_h():
    return _login(ADMIN)


@pytest.fixture(scope="module")
def manager_h():
    return _login(MANAGER)


@pytest.fixture(scope="module")
def driver_h():
    return _login(DRIVER)


@pytest.fixture(scope="module")
def admin_b_h():
    return _login(ADMIN_B)


@pytest.fixture(scope="module")
def mongo():
    mc = pymongo.MongoClient(os.environ["MONGO_URL"], serverSelectionTimeoutMS=5000)
    yield mc[os.environ["DB_NAME"]]
    mc.close()


@pytest.fixture(scope="module", autouse=True)
def seed_and_restore(mongo):
    """Seed trips TEST-WKR + sauvegarde/restauration de la config settings."""
    settings_q = {"tenant_id": "default", "id": "default"}
    original = mongo.settings.find_one(settings_q, {"_id": 0,
                                                    "weekly_report_enabled": 1,
                                                    "weekly_report_recipients": 1}) or {}
    mongo.trips.delete_many({"id": {"$regex": f"^{MARK}-"}})
    now = datetime.utcnow().isoformat() + "+00:00"
    for t in SEED_TRIPS:
        mongo.trips.insert_one({**t, "tenant_id": "default", "created_at": now,
                                "duration_min": 30, "auto_classified": False})
    yield
    mongo.trips.delete_many({"id": {"$regex": f"^{MARK}-"}})
    mongo.weekly_report_log.delete_many({"week_monday": TEST_WEEK})
    restore = {"weekly_report_enabled": original.get("weekly_report_enabled", False),
               "weekly_report_recipients": original.get("weekly_report_recipients", [])}
    mongo.settings.update_one(settings_q, {"$set": restore})


def _xlsx_rows(content: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _find_row(rows, name):
    for r in rows:
        if r and r[0] == name:
            return r
    return None


# ---------------------------------------------------------------------------
# 1. Fonctions pures de période
# ---------------------------------------------------------------------------
def test_period_functions():
    from app.weekly_driver_report import monday_of, previous_week_monday, week_bounds
    assert monday_of(date(2025, 3, 5)) == date(2025, 3, 3)
    assert monday_of(date(2025, 3, 3)) == date(2025, 3, 3)
    assert monday_of(date(2025, 3, 9)) == date(2025, 3, 3)
    # Depuis un lundi : la semaine précédente commence 7 jours avant
    from zoneinfo import ZoneInfo
    ref = datetime(2025, 3, 10, 6, 0, tzinfo=ZoneInfo("Europe/Zurich"))
    assert previous_week_monday(ref) == date(2025, 3, 3)
    # Bornes UTC — CET (mars avant changement d'heure) = UTC+1
    start, end = week_bounds(date(2025, 3, 3))
    assert start == "2025-03-02T23:00:00+00:00"
    assert end == "2025-03-09T23:00:00+00:00"
    # Semaine avec passage à l'heure d'été (dernier dimanche de mars 2025)
    start2, end2 = week_bounds(date(2025, 3, 31))
    assert start2 == "2025-03-30T22:00:00+00:00"  # CEST = UTC+2


# ---------------------------------------------------------------------------
# 2. Configuration — GET/PUT + validation + RBAC
# ---------------------------------------------------------------------------
def test_config_get_defaults(admin_h):
    r = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r.status_code == 200, r.text
    d = r.json()
    assert isinstance(d["enabled"], bool)
    assert isinstance(d["recipients"], list)
    assert d["formats"] == ["pdf", "xlsx"]
    assert "smtp_configured" in d
    assert "lundi à dimanche" in d["period_rule"]
    assert d["send_day"] == "Lundi 06:00 (Europe/Zurich)"
    assert d["default_week"]  # lundi de la semaine précédente


def test_config_put_and_normalisation(admin_h):
    r = requests.put(CONFIG_URL, headers=admin_h, timeout=20, json={
        "enabled": True,
        "recipients": ["  Gestion@Logitrak.CH ", "gestion@logitrak.ch", "rh@logitrak.ch"]})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["enabled"] is True
    # lowercase + dédoublonnage
    assert d["recipients"] == ["gestion@logitrak.ch", "rh@logitrak.ch"]
    # relecture cohérente
    r2 = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r2.json()["recipients"] == ["gestion@logitrak.ch", "rh@logitrak.ch"]


def test_config_validation(admin_h):
    r = requests.put(CONFIG_URL, headers=admin_h, timeout=20,
                     json={"recipients": ["pas-un-email"]})
    assert r.status_code == 400
    r2 = requests.put(CONFIG_URL, headers=admin_h, timeout=20,
                      json={"recipients": [f"u{i}@x.ch" for i in range(21)]})
    assert r2.status_code == 400


def test_config_rbac(manager_h, driver_h):
    assert requests.get(CONFIG_URL, headers=manager_h, timeout=20).status_code == 200
    r = requests.put(CONFIG_URL, headers=manager_h, timeout=20, json={"enabled": True})
    assert r.status_code == 200
    assert requests.get(CONFIG_URL, headers=driver_h, timeout=20).status_code == 403
    assert requests.put(CONFIG_URL, headers=driver_h, timeout=20,
                        json={"enabled": True}).status_code == 403
    assert requests.get(DOWNLOAD_URL, headers=driver_h, timeout=20,
                        params={"week": TEST_WEEK, "format": "pdf"}).status_code == 403


# ---------------------------------------------------------------------------
# 3. Téléchargement manuel — PDF / XLSX / semaine vide / normalisation / isolation
# ---------------------------------------------------------------------------
def test_download_pdf(admin_h):
    r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                     params={"week": TEST_WEEK, "format": "pdf"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert r.content[:4] == b"%PDF"
    assert f"rapport_hebdo_chauffeurs_{TEST_WEEK}.pdf" in r.headers.get("content-disposition", "")


def test_download_xlsx_values(admin_h):
    r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                     params={"week": TEST_WEEK, "format": "xlsx"})
    assert r.status_code == 200, r.text
    rows = _xlsx_rows(r.content)
    alice = _find_row(rows, f"{MARK} Alice")
    assert alice, "Ligne Alice absente du XLSX"
    # Colonnes : Chauffeur, Km PRO, Km PRIVÉ, Km non classifiés, Km total,
    #            Trajets, Trajets non classifiés, Véhicules utilisés
    assert alice[1] == 100.5
    assert alice[2] == 50.25
    assert alice[3] == 10.0
    assert alice[4] == 160.75  # total = TOUS les km (pro + privé + non classifiés)
    assert alice[5] == 3
    assert alice[6] == 1
    assert f"{MARK} 111" in alice[7] and f"{MARK} 222" in alice[7]
    bob = _find_row(rows, f"{MARK} Bob")
    assert bob, "Ligne Bob (trajet limite dimanche 23:30 locale) absente"
    assert bob[1] == 200.0 and bob[5] == 1
    # Le trajet T5 (999 km, lundi suivant 00:30 locale) ne doit PAS être compté
    assert alice[1] == 100.5  # pas 1099.5


def test_download_week_normalized_to_monday(admin_h):
    """Passer un mercredi retourne la même semaine que le lundi."""
    r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                     params={"week": "2025-03-05", "format": "xlsx"})
    assert r.status_code == 200
    assert f"rapport_hebdo_chauffeurs_{TEST_WEEK}.xlsx" in r.headers.get("content-disposition", "")
    alice = _find_row(_xlsx_rows(r.content), f"{MARK} Alice")
    assert alice and alice[4] == 160.75


def test_download_empty_week_is_valid_report(admin_h):
    """Aucune donnée : rapport valide « aucune donnée », jamais une erreur serveur."""
    for fmt in ("pdf", "xlsx"):
        r = requests.get(DOWNLOAD_URL, headers=admin_h, timeout=30,
                         params={"week": "2020-01-06", "format": fmt})
        assert r.status_code == 200, f"{fmt}: {r.status_code}"
    rows = _xlsx_rows(r.content)
    assert any(row and row[0] == "Aucun trajet sur la période" for row in rows)


def test_download_bad_params(admin_h):
    assert requests.get(DOWNLOAD_URL, headers=admin_h, timeout=20,
                        params={"week": "pas-une-date", "format": "pdf"}).status_code == 400
    assert requests.get(DOWNLOAD_URL, headers=admin_h, timeout=20,
                        params={"week": TEST_WEEK, "format": "docx"}).status_code == 422


def test_tenant_isolation(admin_b_h):
    """Tenant B : config indépendante + aucune donnée TEST-WKR du tenant default."""
    r = requests.get(CONFIG_URL, headers=admin_b_h, timeout=20)
    assert r.status_code == 200
    assert r.json()["recipients"] != ["gestion@logitrak.ch", "rh@logitrak.ch"] or \
        r.json()["recipients"] == []  # tenant B ne partage pas la config default
    rx = requests.get(DOWNLOAD_URL, headers=admin_b_h, timeout=30,
                      params={"week": TEST_WEEK, "format": "xlsx"})
    assert rx.status_code == 200
    rows = _xlsx_rows(rx.content)
    assert _find_row(rows, f"{MARK} Alice") is None
    assert _find_row(rows, f"{MARK} Bob") is None


# ---------------------------------------------------------------------------
# 4. Job scheduler — SMTP mocké, idempotence, multi-tenant, statuts
# ---------------------------------------------------------------------------
def _run_job(monkeypatch, mongo, *, smtp_configured, send_ok=True,
             enabled=True, recipients=None):
    """Exécute run_weekly_report_job pour TEST_WEEK avec SMTP entièrement mocké.
    Retourne (results, emails_captés)."""
    import app.emailer as emailer_mod

    if recipients is None:
        recipients = ["gestion@logitrak.ch"]
    mongo.settings.update_one(
        {"tenant_id": "default", "id": "default"},
        {"$set": {"weekly_report_enabled": enabled,
                  "weekly_report_recipients": recipients}})

    sent = []

    async def fake_send(to_list, subject, text, html, attachments):
        sent.append({"to": list(to_list), "subject": subject,
                     "attachments": [(a[0], len(a[1])) for a in attachments]})
        return send_ok

    monkeypatch.setattr(emailer_mod, "is_smtp_configured", lambda: smtp_configured)
    monkeypatch.setattr(emailer_mod, "send_email_with_attachments", fake_send)

    async def runner():
        from app.db import close_db, init_db
        from app.weekly_driver_report import run_weekly_report_job
        init_db()
        try:
            return await run_weekly_report_job(week_monday=date(2025, 3, 3))
        finally:
            close_db()

    results = asyncio.run(runner())
    return results, sent


@pytest.fixture()
def clean_log(mongo):
    mongo.weekly_report_log.delete_many({"week_monday": TEST_WEEK})
    yield
    mongo.weekly_report_log.delete_many({"week_monday": TEST_WEEK})


def test_job_sends_one_email_with_two_attachments(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True)
    assert results.get("default") == "sent"
    assert len(sent) == 1
    mail = sent[0]
    assert mail["to"] == ["gestion@logitrak.ch"]
    assert f"semaine du 03/03/2025 au 09/03/2025" in mail["subject"]
    names = [a[0] for a in mail["attachments"]]
    assert names == [f"rapport_hebdo_chauffeurs_{TEST_WEEK}.pdf",
                     f"rapport_hebdo_chauffeurs_{TEST_WEEK}.xlsx"]
    assert all(a[1] > 1000 for a in mail["attachments"])  # fichiers non vides
    log = mongo.weekly_report_log.find_one({"tenant_id": "default", "week_monday": TEST_WEEK})
    assert log and log["status"] == "sent" and log["error"] is None

    # Idempotence : relance même semaine -> aucun nouvel envoi
    results2, sent2 = _run_job(monkeypatch, mongo, smtp_configured=True)
    assert results2.get("default") == "already_sent"
    assert len(sent2) == 0
    assert mongo.weekly_report_log.count_documents(
        {"tenant_id": "default", "week_monday": TEST_WEEK, "status": "sent"}) == 1


def test_job_smtp_not_configured(monkeypatch, mongo, clean_log):
    """SMTP absent : rapport loggé NOT_SENT, zéro email, le job ne plante pas."""
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=False)
    assert results.get("default") == "not_sent_smtp_not_configured"
    assert sent == []
    log = mongo.weekly_report_log.find_one({"tenant_id": "default", "week_monday": TEST_WEEK})
    assert log["status"] == "not_sent_smtp_not_configured"
    # Relance : pas de statut 'sent' -> nouvelle tentative autorisée (pas de blocage)
    results2, _ = _run_job(monkeypatch, mongo, smtp_configured=False)
    assert results2.get("default") == "not_sent_smtp_not_configured"


def test_job_no_recipients(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, recipients=[])
    assert results.get("default") == "not_sent_no_recipients"
    assert sent == []


def test_job_disabled(monkeypatch, mongo, clean_log):
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, enabled=False)
    assert "default" not in results
    assert sent == []
    assert mongo.weekly_report_log.count_documents(
        {"tenant_id": "default", "week_monday": TEST_WEEK}) == 0


def test_job_smtp_error(monkeypatch, mongo, clean_log):
    """Échec SMTP : statut error propre, le job global continue (pas d'exception)."""
    results, sent = _run_job(monkeypatch, mongo, smtp_configured=True, send_ok=False)
    assert results.get("default") == "error"
    assert len(sent) == 1  # tentative faite
    log = mongo.weekly_report_log.find_one({"tenant_id": "default", "week_monday": TEST_WEEK})
    assert log["status"] == "error" and log["error"]


def test_config_reports_last_send(admin_h, monkeypatch, mongo, clean_log):
    """La config expose le dernier envoi (date + statut) après un run."""
    _run_job(monkeypatch, mongo, smtp_configured=True)
    r = requests.get(CONFIG_URL, headers=admin_h, timeout=20)
    assert r.status_code == 200
    last = r.json()["last_send"]
    assert last and last["week_monday"] == TEST_WEEK and last["status"] == "sent"
