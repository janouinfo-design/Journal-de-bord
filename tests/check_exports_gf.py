"""Inspect BEV guard-rail exports downloaded via the UI (xlsx cell empty, pdf '—')."""
import sys
from openpyxl import load_workbook

xlsx = "/app/.screenshots/export_qa.xlsx"
wb = load_workbook(xlsx)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
header = None
for r in rows:
    if r and any(c and "Carburant" in str(c) for c in r):
        header = r
        break
print("HEADER:", header)
idx = None
if header:
    for i, c in enumerate(header):
        if c and "Carburant" in str(c):
            idx = i
print("fuel col index:", idx)
for r in rows:
    if r and any(c and "QA-" in str(c) for c in r):
        print("ROW:", r)
        if idx is not None:
            print("   fuel cell repr:", repr(r[idx]))

# PDF text
try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None
if PdfReader:
    reader = PdfReader("/app/.screenshots/export_qa.pdf")
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    for line in text.split("\n"):
        if "QA-" in line:
            print("PDFLINE:", repr(line))
    print("PDF contains '0.00':", "0.00" in text)
else:
    print("no pdf lib")
